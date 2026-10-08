"""Read the original XLSB without writing it; join contractual overrides by stable ID."""
import argparse, collections, hashlib, json, re
from datetime import datetime, timezone
from pathlib import Path
from pyxlsb import open_workbook, convert_date

ROOT = Path(__file__).resolve().parents[1]
CLOSED = {'提前贖回', '提前到期', '到期贖回', '提前敲出', '到期', '承接股票'}

def date(v):
    if isinstance(v, (int, float)) and 40000 < v < 60000:
        return convert_date(v).date().isoformat()
    return None

def number(v):
    return float(v) if isinstance(v, (int, float)) else None

def rate(v):
    n = number(v)
    if n is None or n == 999: return None
    return n / 100 if n > 2 else n

def amount(v):
    if isinstance(v, (int, float)): return float(v)
    m = re.fullmatch(r'\s*([\d,.]+)\s*([kKmM]?)\s*', str(v or ''))
    if not m: return None
    return float(m[1].replace(',', '')) * {'':1, 'k':1000, 'm':1000000}[m[2].lower()]

def symbol(v):
    raw = ' '.join(str(v).strip().split()).upper()
    if re.search(r'\b(HK|JP|NA)$', raw) or raw.endswith('.T'): return None
    raw = re.sub(r'(\.OQ|\.N)$','',raw)
    m = re.fullmatch(r'([A-Z][A-Z0-9.-]*)(?:\s+(?:UQ|UW|UN|US|UP))?', raw)
    return m[1] if m else None

def extract(path):
    trades, seen, asofs = [], collections.Counter(), []
    with open_workbook(str(path)) as book:
        for name in book.sheets:
            current, pending = None, []
            with book.get_sheet(name) as sheet:
                for row in sheet.rows():
                    r = {c.c:c.v for c in row if c.v is not None}
                    rowno = row[0].r + 1 if row else 0
                    if rowno == 1 and date(r.get(0)): asofs.append(date(r[0]))
                    if rowno<=3: continue
                    # Blank separator rows delimit product blocks. Metadata may be
                    # at the last row of a block; buffer its preceding fixing rows.
                    if not any(date(r.get(k)) for k in (0,6,7)) and not r.get(8):
                        if pending: raise ValueError(f'{name} 第 {pending[0][0]} 列觀察日無成交資料可歸屬')
                        current=None
                        continue
                    if date(r.get(0)) and r.get(8):
                        client = str(r.get(5, '')).removesuffix('.0')
                        key = '|'.join([name, str(r[8]), client, date(r[0])])
                        seen[key] += 1
                        tid = hashlib.sha256((key+'|'+str(seen[key])).encode()).hexdigest()[:16]
                        issues = []
                        if not r.get(42): issues.append('Excel 未填商品狀態')
                        if not isinstance(r.get(1), str): issues.append('發行商欄不是機構名稱')
                        if seen[key] > 1: issues.append('相同 ISIN、客戶及成交日重複，請核對')
                        current = dict(id=tid, tradeDate=date(r[0]), issuer=str(r.get(1, '')).removesuffix('.0'),
                            clientCode=client, isin=str(r[8]), productType=r.get(10) or '未分類',
                            currency=r.get(11) if isinstance(r.get(11),str) else '未確認', tenorMonths=number(r.get(12)),
                            amount=amount(r.get(43)), amountRaw=r.get(43), coupon=rate(r.get(9)), ko=rate(r.get(2)),
                            koRaw=r.get(2), ki=rate(r.get(3)), strike=rate(r.get(4)), status=r.get(42) or '未填狀態',
                            maturityDate=None, observationDates=[], underlyings=[], issues=issues,
                            source={'sheet':name,'row':rowno}, rule={'verified':False,'mode':'離散全數','comparison':'>=',
                            'koEnabled':False if r.get(2)==999 else (True if rate(r.get(2)) else None),
                            'levelVerified':False, 'scheduleVerified':True,
                            'source':'使用者確認 2026-10-07：999 表示無 KO；現有商品無記憶 KO，Fixing Date 按表列觀察日整理',
                            'note':'達到／超過與期初價調整仍需核對 Termsheet；現有商品以非記憶方式整理。'})
                        for i in range(4):
                            if r.get(13+i): current['underlyings'].append(dict(raw=' '.join(str(r[13+i]).split()),
                                symbol=symbol(r[13+i]), initial=number(r.get(17+i))))
                        if current['amount'] is None: issues.append('認購金額格式待確認')
                        trades.append(current)
                        for n,prev in pending:
                            od,md=date(prev.get(6)),date(prev.get(7))
                            if od and od not in current['observationDates']:current['observationDates'].append(od)
                            if md:current['maturityDate']=md
                        if pending:current['source']['blockStart']=pending[0][0]
                        pending=[]
                    if current is not None:
                        obs, mat = date(r.get(6)), date(r.get(7))
                        if obs and obs not in current['observationDates']: current['observationDates'].append(obs)
                        if mat: current['maturityDate'] = mat
                    else:
                        pending.append((rowno,r))
            if pending:raise ValueError(f'{name} 末尾觀察日缺少成交資料')
    for t in trades:
        t['observationDates'].sort()
        if not t['maturityDate']: t['issues'].append('缺少到期日')
        if not t['observationDates']: t['issues'].append('缺少 Fixing Date／KO 觀察日')
        if t['maturityDate'] and any(d>t['maturityDate'] for d in t['observationDates']):t['issues'].append('表列觀察日晚於到期日，請核對')
    return {'schemaVersion':1,'sourceFile':path.name,'sourceAsOf':max(asofs) if asofs else None,
            'generatedAt':datetime.now(timezone.utc).isoformat(), 'asOf':max(asofs) if asofs else None,
            'prices':{},'priceSource':'富途未復權日線收盤價（預警）','priceCheckedAt':None,
            'expectedPriceDate':None,'trades':trades}

def supplement(data, path):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    byid = {t['id']:t for t in data['trades']}
    def rows(sheet):
        vals=book[sheet].iter_rows(values_only=True)
        headers=next(vals)
        return [dict(zip(headers,r), _row=i) for i,r in enumerate(vals, start=2) if r and r[0]]
    def yes(v): return str(v or '').strip() == '是'
    def iso(v): return v.date().isoformat() if isinstance(v,datetime) else str(v or '')[:10]
    # Transactions not yet written into the source XLSB can be recorded in the
    # companion workbook. Keep the source workbook read-only and deduplicate IDs.
    if '補充成交' in book.sheetnames:
        supplemental = rows('補充成交')
        underlyings = collections.defaultdict(list)
        for r in rows('補充標的') if '補充標的' in book.sheetnames else []:
            underlyings[str(r['成交ID'])].append(dict(raw=' '.join(str(r['Bloomberg代號']).split()),
                symbol=symbol(r['Bloomberg代號']), initial=number(r.get('初始價'))))
        for r in supplemental:
            tid = str(r['成交ID'])
            if tid in byid: raise ValueError(f'補充成交 {tid} 已存在於原始 XLSB，請移除重複補充列')
            trade_date = iso(r['成交日'])
            code = str(r.get('客戶代碼') or '').strip()
            isin = str(r.get('ISIN') or '').strip()
            if not code or not isin: raise ValueError(f'補充成交第 {r.get("成交ID")} 列缺少客戶代碼或 ISIN')
            if any(t['isin'] == isin and t['clientCode'] == code and t['tradeDate'] == trade_date for t in data['trades']):
                raise ValueError(f'補充成交 {isin}／{code}／{trade_date} 與原始 XLSB 可能重複，請核對')
            us = underlyings.get(tid, [])
            if not us: raise ValueError(f'補充成交 {tid} 缺少標的及初始價')
            ko_raw = r.get('Autocall Level')
            t = dict(id=tid, tradeDate=trade_date, issuer=str(r.get('發行商') or ''), clientCode=code,
                isin=isin, productType=str(r.get('產品') or '未分類'), currency=str(r.get('幣別') or '未確認'),
                tenorMonths=number(r.get('年期(月)')), amount=amount(r.get('認購金額')),
                amountRaw=r.get('認購金額'), coupon=rate(r.get('年化票息')),
                ko=rate(ko_raw), koRaw=ko_raw, ki=rate(r.get('Barrier／KI比例')),
                strike=rate(r.get('Strike比例')), status=str(r.get('商品狀態') or '在期'),
                issueDate=iso(r.get('Issue Date')), finalValuationDate=iso(r.get('Final Valuation Date')),
                maturityDate=iso(r.get('Maturity Date')), observationDates=[], underlyings=us, issues=[],
                source={'sheet':'補充成交','row':r['_row'],'pimsCode':str(r.get('PIMS代碼') or ''),
                    'blotterId':str(r.get('BlotterID') or '')},
                rule={'verified':False,'mode':'離散全數','comparison':'>=','koEnabled':None,
                    'levelVerified':False,'scheduleVerified':False,'source':'','note':''})
            data['trades'].append(t); byid[tid] = t
    for r in rows('條款確認'):
        tid = str(r['成交ID'])
        if tid not in byid: raise ValueError('補充表含未知成交ID，請先同步最新來源')
        t=byid[tid]
        t['rule'].update(verified=yes(r.get('條款已核對')),mode=r.get('KO觀察方式') or '未確認',
            koEnabled={'是':True,'否':False}.get(r.get('有KO條款')),
            comparison=r.get('觸發比較') or '>=',levelVerified=yes(r.get('期初價及調整已核對')),
            scheduleVerified=yes(r.get('觀察日已核對')),source=r.get('Termsheet來源') or '',note=r.get('備註') or '')
        if r.get('確認KO比例') is not None: t['ko']=number(r['確認KO比例'])
        if r.get('確認到期日'): t['maturityDate']=iso(r['確認到期日'])
        if r.get('確認商品狀態'): t['status']=r['確認商品狀態']
    dates=collections.defaultdict(list)
    for r in rows('觀察日'):
        tid = str(r['成交ID'])
        if tid not in byid: raise ValueError('觀察日含未知成交ID')
        dates[tid].append(iso(r['觀察日']))
    for tid,ds in dates.items(): byid[tid]['observationDates']=sorted(set(ds))
    if '收盤價' in book.sheetnames:
        for r in rows('收盤價'):
            if str(r.get('未復權收盤價')) not in ('None',''):
                d=iso(r['美國交易日']); s=str(r['美股代號']).upper()
                data['prices'].setdefault(s,{})[d]={'close':float(r['未復權收盤價']), 'source':r.get('來源') or '手動匯入', 'verified':yes(r.get('已核對'))}
    if '補充成交' in book.sheetnames: data['supplementFile'] = path.name
    dates_asof = [t['tradeDate'] for t in data['trades']]
    data['sourceAsOf'] = max([d for d in [data.get('sourceAsOf'), *dates_asof] if d])
    book.close()

def main():
    p=argparse.ArgumentParser()
    p.add_argument('source',type=Path);p.add_argument('--supplement',type=Path)
    p.add_argument('--prices',type=Path);p.add_argument('--output',type=Path,default=ROOT/'private/dataset.json')
    a=p.parse_args(); data=extract(a.source)
    if a.supplement: supplement(data,a.supplement)
    if a.prices and a.prices.exists():
        prices=json.loads(a.prices.read_text())
        for s,ds in prices.get('prices',{}).items():
            for dt,v in ds.items():
                if not data['prices'].get(s,{}).get(dt,{}).get('verified'):
                    data['prices'].setdefault(s,{})[dt]=v
        data.update({k:prices.get(k) for k in ['priceCheckedAt','expectedPriceDate','priceSource','priceError','tradingDays']})
    a.output.parent.mkdir(parents=True,exist_ok=True)
    a.output.write_text(json.dumps(data,ensure_ascii=False,indent=2))
    print(json.dumps({'trades':len(data['trades']),'sourceAsOf':data['sourceAsOf'],'output':str(a.output)},ensure_ascii=False))

if __name__=='__main__': main()
